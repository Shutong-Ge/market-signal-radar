# -*- coding: utf-8 -*-
"""
投研舆情雷达 · 投研晨报组装与导出
=================================
把当日的热点榜、板块分布、深度报告、数据管道口径组装成一份可直接发出去的晨报。
  · 幂等组装：按 (日期 + 热点指纹) 生成 brief_key，重复组装覆盖同一条记录而非追加
  · 双格式导出：Markdown（可直接贴进IM/邮件）+ Word（可直接发给投研团队）
  · 口径透明：晨报末尾固定附「数据与方法口径」，让读者知道数字怎么来的
"""
import os, re, json, sqlite3, hashlib
from datetime import datetime, timedelta, timezone
from collections import Counter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(BASE, "data", "radar.db")
OUT = os.path.join(BASE, "output")
CST = timezone(timedelta(hours=8))

def db():
    c = sqlite3.connect(DB)
    c.execute("""CREATE TABLE IF NOT EXISTS briefs(
        id INTEGER PRIMARY KEY AUTOINCREMENT, brief_key TEXT UNIQUE, brief_date TEXT,
        created_at TEXT, n_hotspots INTEGER, md_path TEXT, docx_path TEXT, payload TEXT)""")
    return c

def gather():
    conn = db()
    LEVELN = {5: "爆点", 4: "高热", 3: "升温中", 2: "一般"}
    cls = [dict(label=r[0], category=r[1], n_articles=r[2], n_media=r[3], heat_score=r[4],
                level=LEVELN.get(r[5], "一般"), factors=json.loads(r[6] or "{}"),
                sectors=json.loads(r[7] or "[]"))
           for r in conn.execute("SELECT label,category,n_articles,n_media,heat_score,importance,"
                                 "heat_factors,sectors FROM clusters ORDER BY heat_score DESC")]
    reps = {r[0]: dict(sections=json.loads(r[1]), citations=json.loads(r[2]), degraded=r[3])
            for r in conn.execute("SELECT cluster_label,sections,citations,degraded FROM reports")}
    dd = conn.execute("SELECT n_in,l1,l2,l3,n_out FROM dedup_runs ORDER BY id DESC LIMIT 1").fetchone()
    col = conn.execute("SELECT SUM(fetched),SUM(inserted) FROM collect_runs").fetchone()
    media_n = conn.execute("SELECT COUNT(DISTINCT media) FROM raw_articles").fetchone()[0]
    conn.close()
    return cls, reps, dict(dedup=dd, collect=col, media_n=media_n)

def build_md(cls, reps, meta, day):
    top = cls[:10]
    cat = Counter(c["category"] for c in cls)
    sec = Counter(s for c in cls for s in c["sectors"])
    L = []
    L.append(f"# 投研舆情晨报 · {day}\n")
    L.append(f"> 覆盖 {meta['media_n']} 家媒体、{meta['collect'][1]:,} 条入库资讯；"
             f"三层去重后 {meta['dedup'][4]:,} 条进入分析；识别热点 {len(cls)} 个"
             f"（爆点 {sum(1 for c in cls if c['level']=='爆点')} 个）。\n")
    L.append("## 一、今日热点榜\n")
    L.append("| # | 热度 | 分级 | 板块类目 | 热点 | 报道量 | 媒体 |")
    L.append("|---|------|------|----------|------|--------|------|")
    for i, c in enumerate(top, 1):
        L.append(f"| {i} | {c['heat_score']:.3f} | {c['level']} | {c['category']} | "
                 f"{c['label'][:34]} | {c['n_articles']} | {c['n_media']} |")
    L.append("\n## 二、板块分布\n")
    L.append("类目声量：" + "；".join(f"{k} {v}个" for k, v in cat.most_common()))
    if sec:
        L.append(f"\n关联板块（东财快讯标签，Top5）：{'、'.join(k for k, _ in sec.most_common(5))}")
    L.append("\n## 三、重点事件深度分析\n")
    for c in top[:5]:
        r = reps.get(c["label"])
        if not r: continue
        L.append(f"### {c['label'][:44]}")
        L.append(f"*热度 {c['heat_score']:.3f}（{c['level']}）· {c['n_articles']} 篇 / "
                 f"{c['n_media']} 家媒体 · {c['category']}"
                 f"{' · 模板降级' if r['degraded'] else ''}*\n")
        for k, v in r["sections"].items():
            if v: L.append(f"**{k}**：{v}\n")
        L.append("<details><summary>引用溯源（混合检索 Top8）</summary>\n")
        for ct in r["citations"][:8]:
            L.append(f"- [{ct['n']}] （{ct['media']}）{ct['title'][:52]} "
                     f"`rrf={ct['rrf']} bm25#{ct['bm25_rank']} vec#{ct['vec_rank']}`")
        L.append("\n</details>\n")
    L.append("## 四、数据与方法口径\n")
    d = meta["dedup"]
    L.append(f"- **采集**：{meta['media_n']} 家媒体公开接口，累计抓取 {meta['collect'][0]:,} 条、"
             f"入库 {meta['collect'][1]:,} 条")
    L.append(f"- **三层去重**：{d[0]:,} → {d[4]:,} 条（精确 {d[1]} / SimHash近似 {d[2]} / 语义 {d[3]}）")
    L.append(f"- **聚类**：向量近邻（余弦≥0.80）+ 并查集连通合并，无预设簇数；"
             f"过滤程式化公告簇与单一媒体簇")
    L.append(f"- **热度**：四因子（覆盖度/爆发速度/持续时长/媒体覆盖）窗口内对数归一；"
             f"社交因子因数据源不可采而权重归一剔除；分级按窗口内分位标定")
    L.append(f"- **深度分析**：BM25 + 向量混合检索、RRF 融合，Top5 限流生成，7 天缓存")
    L.append(f"\n*本晨报由「投研舆情雷达」自动生成，结论供研究参考，不构成投资建议。*")
    return "\n".join(L)

def build_docx(cls, reps, meta, day, path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    NAVY, BLUE, GREY = RGBColor(0x16, 0x32, 0x4F), RGBColor(0x1F, 0x6F, 0xB2), RGBColor(0x8A, 0x93, 0x9E)
    doc = Document()
    doc.styles["Normal"].font.name = "Microsoft YaHei"; doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(f"投研舆情晨报 · {day}", level=0).runs[0].font.color.rgb = NAVY
    p = doc.add_paragraph(f"覆盖 {meta['media_n']} 家媒体 · 入库 {meta['collect'][1]:,} 条 · "
                          f"去重后 {meta['dedup'][4]:,} 条 · 识别热点 {len(cls)} 个")
    p.runs[0].font.size = Pt(9); p.runs[0].font.color.rgb = GREY
    doc.add_heading("一、今日热点榜", level=1).runs[0].font.color.rgb = BLUE
    t = doc.add_table(rows=1, cols=6); t.style = "Light Grid Accent 1"
    for j, h in enumerate(["#", "热度", "分级", "类目", "热点", "报道/媒体"]):
        t.rows[0].cells[j].text = h
    for i, c in enumerate(cls[:10], 1):
        cs = t.add_row().cells
        for j, v in enumerate([i, f"{c['heat_score']:.3f}", c["level"], c["category"],
                               c["label"][:30], f"{c['n_articles']}/{c['n_media']}"]):
            cs[j].text = str(v)
    doc.add_heading("二、重点事件深度分析", level=1).runs[0].font.color.rgb = BLUE
    for c in cls[:5]:
        r = reps.get(c["label"])
        if not r: continue
        doc.add_heading(c["label"][:44], level=2)
        sp = doc.add_paragraph(f"热度 {c['heat_score']:.3f}（{c['level']}）· {c['n_articles']} 篇 / "
                               f"{c['n_media']} 家媒体 · {c['category']}")
        sp.runs[0].font.size = Pt(9); sp.runs[0].font.color.rgb = GREY
        for k, v in r["sections"].items():
            if not v: continue
            par = doc.add_paragraph()
            run = par.add_run(f"{k}："); run.bold = True; run.font.color.rgb = NAVY
            par.add_run(v)
        cp = doc.add_paragraph("引用：" + " / ".join(f"[{ct['n']}]{ct['title'][:26]}"
                                                     for ct in r["citations"][:4]))
        cp.runs[0].font.size = Pt(8.5); cp.runs[0].font.color.rgb = GREY
    doc.add_heading("三、数据与方法口径", level=1).runs[0].font.color.rgb = BLUE
    d = meta["dedup"]
    for line in [f"采集：{meta['media_n']} 家媒体公开接口，累计入库 {meta['collect'][1]:,} 条",
                 f"三层去重：{d[0]:,} → {d[4]:,} 条（精确{d[1]}/SimHash{d[2]}/语义{d[3]}）",
                 "聚类：向量近邻+并查集连通合并，无预设簇数",
                 "热度：四因子窗口内对数归一，社交因子因不可采而归一剔除",
                 "深度分析：BM25+向量混合检索、RRF融合，Top5限流、7天缓存"]:
        doc.add_paragraph(line, style="List Bullet")
    fp = doc.add_paragraph("本晨报由「投研舆情雷达」自动生成，结论供研究参考，不构成投资建议。")
    fp.runs[0].font.size = Pt(8.5); fp.runs[0].font.color.rgb = GREY
    doc.save(path)

def run(verbose=True):
    cls, reps, meta = gather()
    if not cls:
        print("无热点数据，请先运行 cluster.py"); return None
    day = datetime.now(CST).strftime("%Y-%m-%d")
    key = hashlib.md5((day + "|" + "|".join(c["label"] for c in cls[:10])).encode()).hexdigest()[:16]
    md = build_md(cls, reps, meta, day)
    md_path = os.path.join(OUT, f"投研舆情晨报_{day}.md")
    docx_path = os.path.join(OUT, f"投研舆情晨报_{day}.docx")
    open(md_path, "w", encoding="utf-8").write(md)
    build_docx(cls, reps, meta, day, docx_path)
    conn = db()
    prev = conn.execute("SELECT id FROM briefs WHERE brief_key=?", (key,)).fetchone()
    conn.execute("INSERT OR REPLACE INTO briefs(brief_key,brief_date,created_at,n_hotspots,"
                 "md_path,docx_path,payload) VALUES(?,?,?,?,?,?,?)",
                 (key, day, datetime.now(CST).strftime("%Y-%m-%d %H:%M:%S"), len(cls),
                  md_path, docx_path, json.dumps(dict(top=[c["label"] for c in cls[:10]]),
                                                 ensure_ascii=False)))
    conn.commit(); conn.close()
    if verbose:
        print(f"[晨报] {day} | 热点 {len(cls)} 个 | 深度分析 {min(5,len(reps))} 篇 | "
              f"{'幂等覆盖（同指纹已存在）' if prev else '新建'}")
        print(f"  Markdown: {md_path} ({len(md)} 字)")
        print(f"  Word:     {docx_path}")
    return dict(day=day, key=key, n_hotspots=len(cls), md_path=md_path, docx_path=docx_path,
                idempotent_overwrite=bool(prev), md_chars=len(md))

if __name__ == "__main__":
    run()
